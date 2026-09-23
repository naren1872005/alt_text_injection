import io
import os
import shutil
import uuid
import zipfile
import csv
import json
import asyncio
import time
from pathlib import Path
from typing import Optional, Dict, Any, List
from pydantic import BaseModel

from fastapi import FastAPI, UploadFile, File, HTTPException, Form
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware

from extractor import FigureExtractor
from excel_parser import ExcelParser, is_duplicate_marker
from matcher import VisualMatcher
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

def sanitize_excel_val(val: Any) -> Any:
    """Removes invalid XML/control characters not permitted in Excel worksheets."""
    if val is None:
        return ""
    if isinstance(val, str):
        return ILLEGAL_CHARACTERS_RE.sub("", val)
    return val

from openpyxl.drawing.spreadsheet_drawing import (
    SpreadsheetDrawing,
    _check_anchor,
    ChartBase,
    Image as XLDrawingImage,
    OneCellAnchor,
    TwoCellAnchor,
    SHEET_DRAWING_NS
)
from openpyxl.packaging.relationship import Relationship

# High-performance in-memory patch for OpenPyXL:
# Ensures drawing relationships are written with relative paths ('../media/') in a single pass,
# preventing Excel corruption warnings and eliminating redundant zip decompress/recompress passes.
def _fast_spreadsheet_drawing_write(self):
    anchors = []
    for idx, obj in enumerate(self.charts + self.images, 1):
        anchor = _check_anchor(obj)
        if isinstance(obj, ChartBase):
            target_path = obj.path.replace('/xl/', '../') if obj.path.startswith('/xl/') else obj.path
            rel = Relationship(type="chart", Target=target_path)
            anchor.graphicFrame = self._chart_frame(idx)
        elif isinstance(obj, XLDrawingImage):
            target_path = obj.path.replace('/xl/media/', '../media/') if obj.path.startswith('/xl/media/') else obj.path
            rel = Relationship(type="image", Target=target_path)
            child = anchor.pic or (anchor.groupShape and anchor.groupShape.pic)
            if not child:
                anchor.pic = self._picture_frame(idx)
            else:
                child.blipFill.blip.embed = f"rId{idx}"

        anchors.append(anchor)
        self._rels.append(rel)

    for a in anchors:
        if isinstance(a, OneCellAnchor):
            self.oneCellAnchor.append(a)
        elif isinstance(a, TwoCellAnchor):
            self.twoCellAnchor.append(a)
        else:
            self.absoluteAnchor.append(a)

    tree = self.to_tree()
    tree.set('xmlns', SHEET_DRAWING_NS)
    return tree

SpreadsheetDrawing._write = _fast_spreadsheet_drawing_write

def save_openpyxl_workbook(wb, file_path: Path) -> Path:
    """Direct single-pass high-speed workbook save."""
    wb.save(str(file_path))
    return file_path

def save_openpyxl_bytes(wb) -> bytes:
    """Direct single-pass high-speed workbook byte serializer."""
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

import sys
import webbrowser
import threading

# Prevent PyInstaller windowed mode / headless crash (AttributeError: 'NoneType' object has no attribute 'isatty')
class NullWriter:
    def write(self, s): pass
    def flush(self): pass
    def isatty(self): return False

if sys.stdout is None:
    sys.stdout = NullWriter()
if sys.stderr is None:
    sys.stderr = NullWriter()

if getattr(sys, 'frozen', False):
    BASE_DIR = Path(getattr(sys, '_MEIPASS', os.path.dirname(sys.executable)))
    APP_DATA_DIR = Path(os.path.dirname(sys.executable))
else:
    BASE_DIR = Path(__file__).resolve().parent
    APP_DATA_DIR = BASE_DIR

STATIC_DIR = BASE_DIR / "static"
SESSIONS_DIR = APP_DATA_DIR / "sessions"

SESSIONS_DIR.mkdir(exist_ok=True)
STATIC_DIR.mkdir(exist_ok=True)

app = FastAPI(title="PDF Figure Tag & Image Extractor", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.middleware("http")
async def add_no_cache_headers(request, call_next):
    response = await call_next(request)
    path = request.url.path
    if path.endswith((".js", ".css", ".html")) or path == "/":
        response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
    return response

# In-memory session cache, cancellation tracking, and live progress registry
session_cache = {}
cancelled_uploads = set()
upload_progress: Dict[str, Dict[str, Any]] = {}

SESSION_MAX_AGE_SECONDS = 4 * 3600  # 4 Hours (14,400 seconds)

def save_session_to_disk(session_id: str, data: Dict[str, Any]):
    """Persists session data to session_metadata.json to survive server restarts."""
    if not session_id or not data:
        return
    session_path = SESSIONS_DIR / session_id
    if session_path.exists():
        try:
            meta_path = session_path / "session_metadata.json"
            with open(meta_path, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2, default=str)
        except Exception as e:
            print(f"Warning persisting session {session_id} to disk: {e}")

def get_session_data(session_id: Optional[str] = None) -> Optional[Dict[str, Any]]:
    """
    Safely retrieves session data from RAM cache or disk persistence (session_metadata.json).
    If the session_id is outdated, missing, or belongs to a previous session before server restart,
    gracefully recovers from the latest active session on disk to prevent 404 errors.
    """
    # 1. Check RAM cache if valid session_id provided
    if session_id and session_id in session_cache:
        touch_session(session_id)
        return session_cache[session_id]

    # 2. Check disk for this specific session
    if session_id:
        session_path = SESSIONS_DIR / session_id
        meta_path = session_path / "session_metadata.json"
        if meta_path.exists():
            try:
                with open(meta_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    session_cache[session_id] = data
                    return data
            except Exception as e:
                print(f"Error loading session metadata from disk for {session_id}: {e}")

    # 3. Fallback: find any active session folder on disk, sorted by newest mtime
    if SESSIONS_DIR.exists():
        session_folders = [p for p in SESSIONS_DIR.iterdir() if p.is_dir()]
        session_folders.sort(key=lambda p: p.stat().st_mtime if p.exists() else 0, reverse=True)
        for folder in session_folders:
            f_meta = folder / "session_metadata.json"
            if f_meta.exists():
                try:
                    with open(f_meta, "r", encoding="utf-8") as f:
                        data = json.load(f)
                        session_cache[folder.name] = data
                        if session_id:
                            session_cache[session_id] = data
                        return data
                except Exception as e:
                    print(f"Error loading fallback session {folder.name}: {e}")

    return None

def resolve_session_path(session_id: Optional[str]) -> Optional[Path]:
    """Resolves the actual existing directory path for a session ID with fallback recovery."""
    session_data = get_session_data(session_id)
    if session_data:
        actual_id = session_data.get("session_id")
        if actual_id:
            p = SESSIONS_DIR / actual_id
            if p.exists():
                return p
    if session_id:
        p = SESSIONS_DIR / session_id
        if p.exists():
            return p
    return None

def touch_session(session_id: Optional[str]):
    """Updates the last modified timestamp of a session folder to prevent expiration while actively used."""
    if not session_id:
        return
    session_path = SESSIONS_DIR / session_id
    if session_path.exists():
        try:
            os.utime(session_path, None)
        except Exception:
            pass

def purge_previous_sessions(active_session_id: str):
    """
    Purges all previous session folders except active_session_id.
    Ensures that when a user starts working on a new document, previous session files
    are immediately cleared out so only the current active session remains on disk and memory.
    """
    if not SESSIONS_DIR.exists() or not active_session_id:
        return
    for session_folder in list(SESSIONS_DIR.iterdir()):
        if session_folder.is_dir() and session_folder.name != active_session_id:
            try:
                shutil.rmtree(session_folder, ignore_errors=True)
                session_cache.pop(session_folder.name, None)
            except Exception as e:
                print(f"Warning purging old session {session_folder.name}: {e}")

def cleanup_expired_sessions(max_age_seconds: float = SESSION_MAX_AGE_SECONDS) -> Dict[str, Any]:
    """
    Automatically purges session folders from disk and memory if idle for > max_age_seconds (default 4 hours).
    Also removes orphaned entries from session_cache.
    """
    now = time.time()
    cleaned_count = 0
    freed_bytes = 0

    if SESSIONS_DIR.exists():
        for session_folder in list(SESSIONS_DIR.iterdir()):
            if session_folder.is_dir():
                try:
                    mtime = session_folder.stat().st_mtime
                    if (now - mtime) > max_age_seconds:
                        sess_id = session_folder.name
                        for root, _, files in os.walk(session_folder):
                            for f in files:
                                try:
                                    freed_bytes += os.path.getsize(os.path.join(root, f))
                                except Exception:
                                    pass
                        shutil.rmtree(session_folder, ignore_errors=True)
                        session_cache.pop(sess_id, None)
                        cleaned_count += 1
                except Exception as e:
                    print(f"Warning during session cleanup for {session_folder.name}: {e}")

    for sess_id in list(session_cache.keys()):
        session_path = SESSIONS_DIR / sess_id
        if not session_path.exists():
            session_cache.pop(sess_id, None)

    return {
        "cleaned_sessions": cleaned_count,
        "freed_mb": round(freed_bytes / (1024 * 1024), 2),
        "active_sessions": len(session_cache)
    }

@app.on_event("startup")
async def start_session_cleanup_worker():
    """Initial session cleanup sweep on server start + periodic background worker (every 15 min)."""
    cleanup_expired_sessions()

    async def periodic_cleanup():
        while True:
            await asyncio.sleep(900)  # Check every 15 minutes
            try:
                cleanup_expired_sessions()
            except Exception as e:
                print(f"Periodic session cleanup error: {e}")

    asyncio.create_task(periodic_cleanup())

@app.post("/api/cleanup-sessions")
def manual_cleanup_sessions():
    """Endpoint to manually trigger cleanup of sessions older than 4 hours."""
    stats = cleanup_expired_sessions()
    return JSONResponse(content={"status": "ok", **stats})

def update_progress(upload_id: Optional[str], percent: int, title: str = "Processing Excel…", status: str = ""):
    if not upload_id:
        return
    upload_progress[upload_id] = {
        "percent": max(0, min(100, int(percent))),
        "title": title,
        "status": status,
        "cancelled": upload_id in cancelled_uploads
    }

class CancelUploadRequest(BaseModel):
    upload_id: Optional[str] = None
    session_id: Optional[str] = None

@app.get("/api/upload-progress/{upload_id}")
async def get_upload_progress(upload_id: str):
    prog = upload_progress.get(upload_id)
    if not prog:
        return JSONResponse(content={"percent": 0, "title": "Processing…", "status": "Waiting...", "cancelled": False})
    return JSONResponse(content=prog)

def auto_attach_excel_if_available(session_id: str, figures: List[Dict[str, Any]], formulas: Optional[List[Dict[str, Any]]] = None):
    """
    If an Excel file exists in this session's folder,
    parse all records, generate drawing crops, and run VisualMatcher for figures and formulas.
    """
    touch_session(session_id)
    session_path = SESSIONS_DIR / session_id
    excel_candidates = [
        session_path / "manifest.xlsx",
        session_path / "manifest.xlsm"
    ]
    excel_path = None
    for p in excel_candidates:
        if p.exists():
            excel_path = p
            break
            
    if not excel_path:
        return figures, (formulas or []), [], None

    excel_images_dir = session_path / "excel_images"
    excel_images_dir.mkdir(parents=True, exist_ok=True)
    
    try:
        parser = ExcelParser(str(excel_path))
        excel_records = parser.parse_all_records(output_dir=str(excel_images_dir))
        parser.close()
        
        for rec in excel_records:
            if rec.get("image_filename"):
                rec["image_url"] = f"/api/excel-image/{session_id}/{rec['image_filename']}"
                
        pdf_fn = session_cache.get(session_id, {}).get("filename")
        matcher = VisualMatcher(excel_records, str(excel_images_dir), pdf_filename=pdf_fn)
        if figures:
            matched_figures = matcher.match_figures(figures)
            for fig in matched_figures:
                ex = fig.get("excel_match")
                if ex and ex.get("image_filename"):
                    ex["image_url"] = f"/api/excel-image/{session_id}/{ex['image_filename']}"
                fig["excel_match"] = ex
            figures = matched_figures

        if formulas:
            matched_formulas = matcher.match_formulas(formulas)
            for form in matched_formulas:
                ex = form.get("excel_match")
                if ex and ex.get("image_filename"):
                    ex["image_url"] = f"/api/excel-image/{session_id}/{ex['image_filename']}"
                form["excel_match"] = ex
            formulas = matched_formulas
            
        return figures, (formulas or []), excel_records, excel_path.name
    except Exception as e:
        print(f"Excel attach failed: {e}")
        return figures, (formulas or []), [], None

@app.get("/api/health")
def health_check():
    return {"status": "ok", "service": "PDF Figure Tag Extractor"}

@app.post("/api/upload-pdf")
async def upload_pdf(file: UploadFile = File(...), session_id: Optional[str] = Form(None)):
    """
    Receives an uploaded PDF, parses its StructTreeRoot for all /Figure tags,
    computes bounding boxes from MCIDs, and renders high-res crops for each figure.
    Does NOT use default sample files. Works for any user PDF.
    """
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Uploaded file must be a PDF.")

    if not session_id or session_id not in session_cache:
        session_id = str(uuid.uuid4())
        purge_previous_sessions(session_id)
        session_cache[session_id] = {
            "session_id": session_id,
            "figures": [],
            "figures_count": 0
        }
    else:
        touch_session(session_id)

    session_path = SESSIONS_DIR / session_id
    session_path.mkdir(parents=True, exist_ok=True)

    input_pdf_path = session_path / "input.pdf"
    figures_output_dir = session_path / "figures"
    figures_output_dir.mkdir(parents=True, exist_ok=True)

    # Save uploaded file
    contents = await file.read()
    with open(input_pdf_path, "wb") as f:
        f.write(contents)

    # Run extraction
    extractor = None
    try:
        extractor = FigureExtractor(str(input_pdf_path))
        raw_figures = extractor.extract_figures_with_crops(str(figures_output_dir), dpi=150)
        
        # Enrich figures with image URLs
        figures = []
        for fig in raw_figures:
            crop_fn = fig.get("crop_filename")
            image_url = f"/api/figure-image/{session_id}/{crop_fn}" if crop_fn else None
            figures.append({
                **fig,
                "image_url": image_url
            })

        # Extract formulas with high-res crops
        formulas_output_dir = session_path / "formulas"
        formulas_output_dir.mkdir(parents=True, exist_ok=True)
        raw_formulas = extractor.extract_formulas_with_crops(str(formulas_output_dir), dpi=150)
        formulas = []
        for form in raw_formulas:
            crop_fn = form.get("crop_filename")
            image_url = f"/api/formula-image/{session_id}/{crop_fn}" if crop_fn else None
            formulas.append({
                **form,
                "image_url": image_url
            })

        # Attach Excel manifest if an Excel file was uploaded to this session
        figures, formulas, excel_records, excel_fn = auto_attach_excel_if_available(session_id, figures, formulas)

        has_alt_count = sum(1 for f in figures if f["has_alt"] or f.get("excel_match"))
        missing_alt_count = len(figures) - has_alt_count

        has_formula_alt_count = sum(1 for f in formulas if f["has_alt"] or f.get("excel_match"))
        missing_formula_alt_count = len(formulas) - has_formula_alt_count

        excel_images_count = sum(1 for r in excel_records if r.get("has_image"))
        excel_has_alt = sum(1 for r in excel_records if r.get("has_alt"))

        result = {
            "session_id": session_id,
            "filename": file.filename,
            "total_pages": extractor.page_count,
            "has_struct_tree": extractor.has_struct_tree,
            "figures_count": len(figures),
            "has_alt_count": has_alt_count,
            "missing_alt_count": missing_alt_count,
            "figures": figures,
            "formulas_count": len(formulas),
            "has_formula_alt_count": has_formula_alt_count,
            "missing_formula_alt_count": missing_formula_alt_count,
            "formulas": formulas,
            "excel_filename": excel_fn,
            "excel_records": excel_records,
            "excel_total_records": len(excel_records),
            "excel_images_count": excel_images_count,
            "excel_has_alt_count": excel_has_alt,
            "matched_figures": figures,
            "matched_formulas": formulas
        }

        # Cache session data
        session_cache[session_id] = result
        save_session_to_disk(session_id, result)
        return JSONResponse(content=result)

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error processing PDF: {str(e)}")
    finally:
        if extractor:
            extractor.close()

@app.post("/api/load-sample")
async def load_sample():
    """
    Loads local sample PDF if available in project directory.
    """
    sample_candidates = list(BASE_DIR.glob("*.pdf"))
    sample_path = sample_candidates[0] if sample_candidates else None

    if not sample_path:
        raise HTTPException(status_code=404, detail="No sample PDF file found in project directory. Please upload a PDF.")

    session_id = str(uuid.uuid4())
    purge_previous_sessions(session_id)
    session_path = SESSIONS_DIR / session_id
    session_path.mkdir(parents=True, exist_ok=True)
    pdf_path = session_path / "input.pdf"

    shutil.copyfile(sample_path, pdf_path)

    extractor = None
    try:
        extractor = FigureExtractor(str(pdf_path))
        figures_output_dir = session_path / "figures"
        figures_output_dir.mkdir(parents=True, exist_ok=True)
        raw_figures = extractor.extract_figures_with_crops(str(figures_output_dir), dpi=150)
        
        figures = []
        for fig in raw_figures:
            crop_fn = fig.get("crop_filename")
            image_url = f"/api/figure-image/{session_id}/{crop_fn}" if crop_fn else None
            figures.append({
                **fig,
                "image_url": image_url
            })

        # Extract formulas with high-res crops
        formulas_output_dir = session_path / "formulas"
        formulas_output_dir.mkdir(parents=True, exist_ok=True)
        raw_formulas = extractor.extract_formulas_with_crops(str(formulas_output_dir), dpi=150)
        formulas = []
        for form in raw_formulas:
            crop_fn = form.get("crop_filename")
            image_url = f"/api/formula-image/{session_id}/{crop_fn}" if crop_fn else None
            formulas.append({
                **form,
                "image_url": image_url
            })

        # Auto-attach Excel manifest if present in session
        figures, formulas, excel_records, excel_fn = auto_attach_excel_if_available(session_id, figures, formulas)

        has_alt_count = sum(1 for f in figures if f["has_alt"] or f.get("excel_match"))
        missing_alt_count = len(figures) - has_alt_count

        has_formula_alt_count = sum(1 for f in formulas if f["has_alt"] or f.get("excel_match"))
        missing_formula_alt_count = len(formulas) - has_formula_alt_count

        excel_images_count = sum(1 for r in excel_records if r.get("has_image"))
        excel_has_alt = sum(1 for r in excel_records if r.get("has_alt"))

        result = {
            "session_id": session_id,
            "filename": sample_path.name,
            "total_pages": extractor.page_count,
            "has_struct_tree": extractor.has_struct_tree,
            "figures_count": len(figures),
            "has_alt_count": has_alt_count,
            "missing_alt_count": missing_alt_count,
            "figures": figures,
            "formulas_count": len(formulas),
            "has_formula_alt_count": has_formula_alt_count,
            "missing_formula_alt_count": missing_formula_alt_count,
            "formulas": formulas,
            "excel_filename": excel_fn,
            "excel_records": excel_records,
            "excel_total_records": len(excel_records),
            "excel_images_count": excel_images_count,
            "excel_has_alt_count": excel_has_alt,
            "matched_figures": figures,
            "matched_formulas": formulas
        }

        session_cache[session_id] = result
        save_session_to_disk(session_id, result)
        return JSONResponse(content=result)

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error processing sample PDF: {str(e)}")
    finally:
        if extractor:
            extractor.close()

@app.post("/api/cancel-excel-upload")
async def cancel_excel_upload(req: CancelUploadRequest):
    """
    Cancels an active Excel manifest upload or processing task,
    halts any ongoing extraction, and cleans up temporary manifest files/images.
    """
    if req.upload_id:
        cancelled_uploads.add(req.upload_id)

    if req.session_id and req.session_id in session_cache:
        session_path = SESSIONS_DIR / req.session_id
        # Remove manifest file
        for m_name in ["manifest.xlsx", "manifest.xlsm"]:
            manifest_file = session_path / m_name
            if manifest_file.exists():
                try:
                    manifest_file.unlink()
                except Exception:
                    pass
        # Remove partial extracted excel images
        excel_images_dir = session_path / "excel_images"
        if excel_images_dir.exists():
            try:
                shutil.rmtree(excel_images_dir, ignore_errors=True)
            except Exception:
                pass
        # Remove excel fields from session_cache if they were partially added
        session_cache[req.session_id].pop("excel_records", None)
        session_cache[req.session_id].pop("excel_filename", None)
        session_cache[req.session_id].pop("excel_total_records", None)
        session_cache[req.session_id].pop("excel_images_count", None)
        session_cache[req.session_id].pop("excel_has_alt_count", None)
        session_cache[req.session_id].pop("excel_missing_alt_count", None)

    return JSONResponse(content={"status": "cancelled", "upload_id": req.upload_id})

@app.post("/api/upload-excel")
async def upload_excel(
    file: UploadFile = File(...),
    session_id: Optional[str] = Form(None),
    upload_id: Optional[str] = Form(None)
):
    """
    Upload an Excel ALT Manifest (.xlsx), extract all image records and authoritative ALT texts,
    and match them visually with the current PDF figures and math formulas.
    """
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Uploaded file must be an Excel workbook (.xlsx).")

    is_cancelled = lambda: (upload_id is not None and upload_id in cancelled_uploads)

    if is_cancelled():
        if upload_id in cancelled_uploads:
            cancelled_uploads.remove(upload_id)
        return JSONResponse(status_code=499, content={"detail": "Upload cancelled by user", "cancelled": True})

    session_data = get_session_data(session_id)
    if not session_id or not session_data:
        # If no active session, create one
        session_id = str(uuid.uuid4())
        purge_previous_sessions(session_id)
        session_cache[session_id] = {
            "session_id": session_id,
            "figures": [],
            "figures_count": 0
        }
    else:
        actual_session_id = session_data.get("session_id", session_id)
        session_id = actual_session_id
        session_cache[session_id] = session_data
        touch_session(session_id)

    session_path = SESSIONS_DIR / session_id
    session_path.mkdir(parents=True, exist_ok=True)
    excel_path = session_path / "manifest.xlsx"
    excel_images_dir = session_path / "excel_images"
    excel_images_dir.mkdir(parents=True, exist_ok=True)

    contents = await file.read()
    if is_cancelled():
        if upload_id in cancelled_uploads:
            cancelled_uploads.remove(upload_id)
        return JSONResponse(status_code=499, content={"detail": "Upload cancelled by user", "cancelled": True})

    update_progress(upload_id, 5, "Processing Excel…", "Saving manifest and parsing workbook structure...")

    with open(excel_path, "wb") as f:
        f.write(contents)

    try:
        def _do_excel_processing():
            if is_cancelled():
                raise RuntimeError("Excel upload cancelled by user")

            def on_init_progress(pct, msg):
                update_progress(upload_id, pct, "Processing Excel…", msg)

            parser = ExcelParser(
                str(excel_path),
                is_cancelled=is_cancelled,
                progress_callback=on_init_progress
            )

            if is_cancelled():
                raise RuntimeError("Excel upload cancelled by user")

            # Extract images and records with cancellation and progress support
            def on_excel_parse_progress(pct, msg):
                # Scale 40% to 70%
                current_pct = int(40 + (pct * 0.30))
                update_progress(upload_id, current_pct, "Processing Excel…", msg)

            excel_records = parser.parse_all_records(
                output_dir=str(excel_images_dir),
                is_cancelled=is_cancelled,
                progress_callback=on_excel_parse_progress
            )
            parser.close()

            if is_cancelled():
                raise RuntimeError("Excel upload cancelled by user")

            update_progress(upload_id, 70, "Processing Excel…", f"Extracted {len(excel_records)} manifest records. Linking previews...")

            # Add image URLs
            for rec in excel_records:
                if rec.get("image_filename"):
                    rec["image_url"] = f"/api/excel-image/{session_id}/{rec['image_filename']}"

            # Run visual matcher if session already has PDF figures or formulas
            pdf_figures = session_cache[session_id].get("figures", [])
            pdf_formulas = session_cache[session_id].get("formulas", [])
            if pdf_figures or pdf_formulas:
                pdf_fn = session_cache[session_id].get("filename")
                matcher = VisualMatcher(excel_records, str(excel_images_dir), pdf_filename=pdf_fn)

                def on_match_fig_progress(pct, msg):
                    current_pct = int(70 + (pct * 0.15))
                    update_progress(upload_id, current_pct, "Matching Figures…", msg)

                def on_match_form_progress(pct, msg):
                    current_pct = int(85 + (pct * 0.14))
                    update_progress(upload_id, current_pct, "Matching Formulas…", msg)

                if pdf_figures:
                    matched_figures = matcher.match_figures(
                        pdf_figures,
                        is_cancelled=is_cancelled,
                        progress_callback=on_match_fig_progress
                    )
                    for fig in matched_figures:
                        ex = fig.get("excel_match")
                        if ex and ex.get("image_filename"):
                            ex["image_url"] = f"/api/excel-image/{session_id}/{ex['image_filename']}"
                        fig["excel_match"] = ex
                    session_cache[session_id]["figures"] = matched_figures
                if pdf_formulas:
                    matched_formulas = matcher.match_formulas(
                        pdf_formulas,
                        is_cancelled=is_cancelled,
                        progress_callback=on_match_form_progress
                    )
                    for form in matched_formulas:
                        ex = form.get("excel_match")
                        if ex and ex.get("image_filename"):
                            ex["image_url"] = f"/api/excel-image/{session_id}/{ex['image_filename']}"
                        form["excel_match"] = ex
                    session_cache[session_id]["formulas"] = matched_formulas

            update_progress(upload_id, 100, "Processing Complete", "Finalizing manifest gallery...")
            return excel_records

        excel_records = await asyncio.to_thread(_do_excel_processing)

        # Calculate metrics for Excel records
        excel_images_count = sum(1 for r in excel_records if r.get("has_image"))
        excel_has_alt = sum(1 for r in excel_records if r.get("has_alt"))
        excel_missing_alt = len(excel_records) - excel_has_alt

        # Calculate metrics for PDF Figures
        matched_figs = session_cache[session_id].get("figures", [])
        has_alt_count = sum(1 for f in matched_figs if f.get("has_alt") or f.get("alt_text") or (f.get("excel_match") and f["excel_match"].get("alt_text")))
        missing_alt_count = len(matched_figs) - has_alt_count

        # Calculate metrics for PDF Formulas
        matched_forms = session_cache[session_id].get("formulas", [])
        has_formula_alt_count = sum(1 for f in matched_forms if f.get("has_alt") or f.get("alt_text") or (f.get("excel_match") and f["excel_match"].get("alt_text")) or f.get("actual_text"))
        missing_formula_alt_count = len(matched_forms) - has_formula_alt_count

        session_cache[session_id]["excel_records"] = excel_records
        session_cache[session_id]["excel_filename"] = file.filename
        session_cache[session_id]["excel_total_records"] = len(excel_records)
        session_cache[session_id]["excel_images_count"] = excel_images_count
        session_cache[session_id]["excel_has_alt_count"] = excel_has_alt
        session_cache[session_id]["excel_missing_alt_count"] = excel_missing_alt
        session_cache[session_id]["has_alt_count"] = has_alt_count
        session_cache[session_id]["missing_alt_count"] = missing_alt_count
        session_cache[session_id]["has_formula_alt_count"] = has_formula_alt_count
        session_cache[session_id]["missing_formula_alt_count"] = missing_formula_alt_count

        save_session_to_disk(session_id, session_cache[session_id])

        return JSONResponse(content={
            "session_id": session_id,
            "excel_filename": file.filename,
            "excel_total_records": len(excel_records),
            "excel_images_count": excel_images_count,
            "excel_has_alt_count": excel_has_alt,
            "excel_missing_alt_count": excel_missing_alt,
            "matched_figures": matched_figs,
            "figures_count": len(matched_figs),
            "has_alt_count": has_alt_count,
            "missing_alt_count": missing_alt_count,
            "matched_formulas": matched_forms,
            "formulas_count": len(matched_forms),
            "has_formula_alt_count": has_formula_alt_count,
            "missing_formula_alt_count": missing_formula_alt_count,
            "excel_records": excel_records
        })

    except RuntimeError as re:
        if "cancelled" in str(re).lower():
            # Clean up files on cancellation
            if excel_path.exists():
                try:
                    excel_path.unlink()
                except Exception:
                    pass
            if excel_images_dir.exists():
                try:
                    shutil.rmtree(excel_images_dir, ignore_errors=True)
                except Exception:
                    pass
            return JSONResponse(status_code=499, content={"detail": "Upload cancelled by user", "cancelled": True})
        raise HTTPException(status_code=500, detail=f"Error processing Excel manifest: {str(re)}")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error processing Excel manifest: {str(e)}")
    finally:
        if upload_id and upload_id in cancelled_uploads:
            cancelled_uploads.remove(upload_id)

@app.post("/api/load-sample-excel")
async def load_sample_excel(session_id: Optional[str] = None):
    """
    Loads local sample Excel if available in project directory.
    """
    sample_candidates = list(BASE_DIR.glob("*.xlsx")) + list(BASE_DIR.glob("*.xlsm"))
    sample_path = sample_candidates[0] if sample_candidates else None

    if not sample_path:
        raise HTTPException(status_code=404, detail="No sample Excel manifest found in project directory. Please upload an Excel manifest.")

    session_data = get_session_data(session_id)
    if not session_data:
        session_id = str(uuid.uuid4())
        session_cache[session_id] = {
            "session_id": session_id,
            "figures": [],
            "figures_count": 0,
            "formulas": [],
            "formulas_count": 0
        }
    else:
        actual_session_id = session_data.get("session_id", session_id)
        session_id = actual_session_id
        session_cache[session_id] = session_data
        touch_session(session_id)

    session_path = SESSIONS_DIR / session_id
    session_path.mkdir(parents=True, exist_ok=True)
    excel_path = session_path / "manifest.xlsx"
    excel_images_dir = session_path / "excel_images"
    excel_images_dir.mkdir(parents=True, exist_ok=True)

    if not excel_path.exists():
        shutil.copyfile(sample_path, excel_path)

    try:
        parser = ExcelParser(str(excel_path))
        excel_records = parser.parse_all_records(output_dir=str(excel_images_dir))
        parser.close()

        for rec in excel_records:
            if rec.get("image_filename"):
                rec["image_url"] = f"/api/excel-image/{session_id}/{rec['image_filename']}"

        pdf_figures = session_cache[session_id].get("figures", [])
        pdf_formulas = session_cache[session_id].get("formulas", [])
        if pdf_figures or pdf_formulas:
            matcher = VisualMatcher(excel_records, str(excel_images_dir))
            if pdf_figures:
                matched_figures = matcher.match_figures(pdf_figures)
                for fig in matched_figures:
                    ex = fig.get("excel_match")
                    if ex and ex.get("image_filename"):
                        ex["image_url"] = f"/api/excel-image/{session_id}/{ex['image_filename']}"
                    fig["excel_match"] = ex
                session_cache[session_id]["figures"] = matched_figures
            if pdf_formulas:
                matched_formulas = matcher.match_formulas(pdf_formulas)
                for form in matched_formulas:
                    ex = form.get("excel_match")
                    if ex and ex.get("image_filename"):
                        ex["image_url"] = f"/api/excel-image/{session_id}/{ex['image_filename']}"
                    form["excel_match"] = ex
                session_cache[session_id]["formulas"] = matched_formulas

        excel_images_count = sum(1 for r in excel_records if r.get("has_image"))
        excel_has_alt = sum(1 for r in excel_records if r.get("has_alt"))
        excel_missing_alt = len(excel_records) - excel_has_alt

        # Calculate metrics for PDF Figures
        matched_figs = session_cache[session_id].get("figures", [])
        has_alt_count = sum(1 for f in matched_figs if f.get("has_alt") or f.get("alt_text") or (f.get("excel_match") and f["excel_match"].get("alt_text")))
        missing_alt_count = len(matched_figs) - has_alt_count

        # Calculate metrics for PDF Formulas
        matched_forms = session_cache[session_id].get("formulas", [])
        has_formula_alt_count = sum(1 for f in matched_forms if f.get("has_alt") or f.get("alt_text") or (f.get("excel_match") and f["excel_match"].get("alt_text")) or f.get("actual_text"))
        missing_formula_alt_count = len(matched_forms) - has_formula_alt_count

        session_cache[session_id]["excel_records"] = excel_records
        session_cache[session_id]["excel_filename"] = sample_path.name
        session_cache[session_id]["excel_total_records"] = len(excel_records)
        session_cache[session_id]["excel_images_count"] = excel_images_count
        session_cache[session_id]["excel_has_alt_count"] = excel_has_alt
        session_cache[session_id]["excel_missing_alt_count"] = excel_missing_alt
        session_cache[session_id]["has_alt_count"] = has_alt_count
        session_cache[session_id]["missing_alt_count"] = missing_alt_count
        session_cache[session_id]["has_formula_alt_count"] = has_formula_alt_count
        session_cache[session_id]["missing_formula_alt_count"] = missing_formula_alt_count

        save_session_to_disk(session_id, session_cache[session_id])

        return JSONResponse(content={
            "session_id": session_id,
            "excel_filename": sample_path.name,
            "excel_total_records": len(excel_records),
            "excel_images_count": excel_images_count,
            "excel_has_alt_count": excel_has_alt,
            "excel_missing_alt_count": excel_missing_alt,
            "matched_figures": matched_figs,
            "figures_count": len(matched_figs),
            "has_alt_count": has_alt_count,
            "missing_alt_count": missing_alt_count,
            "matched_formulas": matched_forms,
            "formulas_count": len(matched_forms),
            "has_formula_alt_count": has_formula_alt_count,
            "missing_formula_alt_count": missing_formula_alt_count,
            "excel_records": excel_records
        })

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error processing sample Excel: {str(e)}")

@app.get("/api/excel-image/{session_id}/{filename}")
def get_excel_image(session_id: str, filename: str):
    spath = resolve_session_path(session_id)
    if not spath:
        raise HTTPException(status_code=404, detail="Session not found")
    image_path = spath / "excel_images" / filename
    if not image_path.exists():
        raise HTTPException(status_code=404, detail="Excel image not found")
    return FileResponse(image_path, media_type="image/png", headers={"Cache-Control": "public, max-age=86400"})

@app.get("/api/excel-records/{session_id}")
def get_excel_records(session_id: str, search: Optional[str] = None, page: int = 1, page_size: int = 50):
    session_data = get_session_data(session_id)
    if not session_data or "excel_records" not in session_data:
        raise HTTPException(status_code=404, detail="No Excel records in session")
    records = session_data["excel_records"]
    if search:
        s = search.lower()
        records = [r for r in records if s in (r.get("filename") or "").lower() or s in (r.get("alt_text") or "").lower()]
    total = len(records)
    start = (page - 1) * page_size
    end = start + page_size
    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "records": records[start:end]
    }

@app.get("/api/figure-image/{session_id}/{filename}")
def get_figure_image(session_id: str, filename: str):
    spath = resolve_session_path(session_id)
    if not spath:
        raise HTTPException(status_code=404, detail="Session not found")
    image_path = spath / "figures" / filename
    if not image_path.exists():
        raise HTTPException(status_code=404, detail="Image not found")
    return FileResponse(image_path, media_type="image/png", headers={"Cache-Control": "public, max-age=86400"})

@app.get("/api/formula-image/{session_id}/{filename}")
def get_formula_image(session_id: str, filename: str):
    spath = resolve_session_path(session_id)
    if not spath:
        raise HTTPException(status_code=404, detail="Session not found")
    image_path = spath / "formulas" / filename
    if not image_path.exists():
        raise HTTPException(status_code=404, detail="Formula image not found")
    return FileResponse(image_path, media_type="image/png", headers={"Cache-Control": "public, max-age=86400"})

@app.get("/api/download-formulas-zip/{session_id}")
def download_formulas_zip(session_id: str):
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session formulas not found")

    actual_session_id = session_data.get("session_id", session_id)
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    formulas_dir = session_path / "formulas"
    if not formulas_dir.exists():
        raise HTTPException(status_code=404, detail="Session formulas not found")

    zip_path = session_path / f"pdf_formulas_{actual_session_id[:8]}.zip"
    csv_buffer = io.StringIO()
    writer = csv.writer(csv_buffer)
    writer.writerow(["Formula ID", "Page", "MCIDs", "Status", "Alt Text", "Actual Text", "BBox [x0, y0, x1, y1]", "Width (pt)", "Height (pt)", "Image File"])

    formulas_meta = session_data.get("formulas", [])
    meta_by_file = {f.get("crop_filename"): f for f in formulas_meta}

    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for img_file in sorted(formulas_dir.glob("*.png")):
            zip_file.write(img_file, arcname=f"formulas/{img_file.name}")
            meta = meta_by_file.get(img_file.name, {})
            writer.writerow([
                meta.get("formula_id", ""),
                meta.get("page_number", ""),
                str(meta.get("mcids", "")),
                meta.get("status", ""),
                meta.get("alt_text", "") or "",
                meta.get("actual_text", "") or "",
                str(meta.get("bbox", "")),
                meta.get("bbox_width", ""),
                meta.get("bbox_height", ""),
                img_file.name
            ])

        zip_file.writestr("formulas_manifest.csv", csv_buffer.getvalue())

    return FileResponse(
        zip_path,
        media_type="application/zip",
        filename=f"extracted_formulas_{actual_session_id[:8]}.zip"
    )

@app.get("/api/download-excel-zip/{session_id}")
def download_excel_zip(session_id: str):
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Excel images not found for session")

    actual_session_id = session_data.get("session_id", session_id)
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    excel_dir = session_path / "excel_images"
    if not excel_dir.exists():
        raise HTTPException(status_code=404, detail="Excel images not found for session")

    zip_path = session_path / f"excel_images_{actual_session_id[:8]}.zip"
    csv_buffer = io.StringIO()
    writer = csv.writer(csv_buffer)
    writer.writerow(["Row", "Sr No", "Filename", "Has Image", "Has Alt", "Authoritative Alt Text", "Original Alt", "Updated Alt"])

    records = session_data.get("excel_records", [])
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for rec in records:
            if rec.get("image_filename"):
                img_path = excel_dir / rec["image_filename"]
                if img_path.exists():
                    zip_file.write(img_path, arcname=f"excel_images/{rec['image_filename']}")
            writer.writerow([
                rec.get("row_index", ""),
                rec.get("sr_no", ""),
                rec.get("filename", ""),
                rec.get("has_image", False),
                rec.get("has_alt", False),
                rec.get("alt_text", ""),
                rec.get("original_alt", ""),
                rec.get("updated_alt", "")
            ])

        zip_file.writestr("excel_manifest.csv", csv_buffer.getvalue())

    return FileResponse(
        zip_path,
        media_type="application/zip",
        filename=f"extracted_excel_media_{session_id[:8]}.zip"
    )

def populate_existing_alt_sheet(ws, items, target_dir, label_prefix="PDF Image"):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    alt_header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    alt_header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        "S.No",
        "Page No",
        f"{label_prefix} Preview",
        "Status",
        "Existing Injected Alt Text"
    ]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = alt_header_fill if col_num == len(headers) else header_fill
        cell.font = alt_header_font if col_num == len(headers) else header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"

    row_idx = 2
    for seq_idx, item in enumerate(items, start=1):
        pdf_img_name = item.get("crop_filename", "")
        page_no = item.get("page_number", "")
        existing_alt = (item.get("alt_text") or item.get("actual_text") or (item.get("excel_match") and item["excel_match"].get("alt_text")) or "").strip()
        has_alt = bool(item.get("has_alt") or existing_alt)
        status_label = item.get("status_label") or ("Has Alt Text" if has_alt and existing_alt else "Missing Alt Text")

        ws.append([
            sanitize_excel_val(seq_idx),
            sanitize_excel_val(page_no),
            "", # PDF Image Placeholder (Col C)
            sanitize_excel_val(status_label), # Status (Col D)
            sanitize_excel_val(existing_alt)  # Existing Injected Alt Text (Col E)
        ])

        thumb_h = 0
        if pdf_img_name:
            pdf_img_path = target_dir / pdf_img_name
            if pdf_img_path.exists():
                try:
                    with PILImage.open(pdf_img_path) as pimg:
                        orig_w, orig_h = pimg.size
                        max_w, max_h = 160, 100
                        scale = min(max_w / orig_w, max_h / orig_h, 1.0)
                        thumb_w = int(orig_w * scale)
                        thumb_h = int(orig_h * scale)

                    xl_img = XLImage(str(pdf_img_path))
                    xl_img.width = thumb_w
                    xl_img.height = thumb_h
                    ws.add_image(xl_img, f"C{row_idx}")
                except Exception as img_err:
                    print(f"Warning embedding {label_prefix} image {pdf_img_name}: {img_err}")
                    ws.cell(row=row_idx, column=3).value = pdf_img_name

        row_height = max(80, int(thumb_h * 0.75) + 15) if thumb_h > 0 else 40

        for c_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=c_idx)
            c.border = thin_border
            if c_idx in [1, 2, 3, 4]:
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                if existing_alt:
                    c.fill = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")
        ws.row_dimensions[row_idx].height = row_height
        row_idx += 1

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 32
    ws.column_dimensions['D'].width = 24
    ws.column_dimensions['E'].width = 75


def populate_missing_alt_sheet(ws, items, target_dir, excel_dir, records_by_row, label_prefix="PDF Image"):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    client_col_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    client_col_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        "S.No",
        "Page No",
        f"{label_prefix} Preview",
        "Excel Candidate Image",
        "Excel Alt Text",
        "Status",
        "Client Required Alt Text"
    ]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = client_col_fill if col_num == len(headers) else header_fill
        cell.font = client_col_font if col_num == len(headers) else header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"

    row_idx = 2
    for seq_idx, item in enumerate(items, start=1):
        pdf_img_name = item.get("crop_filename", "")
        page_no = item.get("page_number", "")
        
        # Check for suggested candidate match from Excel below threshold
        cand_match = item.get("candidate_match") or item.get("excel_match")
        best_row = None
        if not cand_match and item.get("match_debug") and item["match_debug"].get("best_single_row"):
            best_row = item["match_debug"]["best_single_row"]
            cand_match = records_by_row.get(best_row)
        elif cand_match:
            best_row = cand_match.get("row") or cand_match.get("row_index")

        cand_img_name = cand_match.get("image_filename") if cand_match else None
        cand_alt_val = (cand_match.get("alt_text") or cand_match.get("raw_alt_text") or "") if cand_match else ""
        conf_val = item.get("confidence") or (item.get("match_debug") and item["match_debug"].get("score")) or 0.0
        conf_pct = int(float(conf_val) * 100) if conf_val else 0

        if cand_match and best_row:
            status_val = f"Low Match - Verify Row {best_row} ({conf_pct}% match)"
        else:
            status_val = "No Match in Excel"

        # Client Required Alt Text is intentionally kept blank for client input
        ws.append([
            sanitize_excel_val(seq_idx),
            sanitize_excel_val(page_no),
            "", # PDF Image Placeholder (Col C)
            "" if cand_img_name else "No Candidate Image", # Excel Image Placeholder (Col D)
            sanitize_excel_val(cand_alt_val), # Excel Alt Text (Col E)
            sanitize_excel_val(status_val),   # Status (Col F)
            ""  # Client Required Alt Text (Col G - Always empty for client to type)
        ])

        thumb_h_pdf = 0
        thumb_h_excel = 0

        # Embed PDF Crop Image in Column C
        if pdf_img_name:
            pdf_img_path = target_dir / pdf_img_name
            if pdf_img_path.exists():
                try:
                    with PILImage.open(pdf_img_path) as pimg:
                        orig_w, orig_h = pimg.size
                        max_w, max_h = 145, 95
                        scale = min(max_w / orig_w, max_h / orig_h, 1.0)
                        thumb_w_pdf = int(orig_w * scale)
                        thumb_h_pdf = int(orig_h * scale)

                    xl_img_pdf = XLImage(str(pdf_img_path))
                    xl_img_pdf.width = thumb_w_pdf
                    xl_img_pdf.height = thumb_h_pdf
                    ws.add_image(xl_img_pdf, f"C{row_idx}")
                except Exception as img_err:
                    print(f"Warning embedding PDF image {pdf_img_name}: {img_err}")
                    ws.cell(row=row_idx, column=3).value = pdf_img_name

        # Embed Excel Manifest Candidate Image in Column D
        if cand_img_name and excel_dir.exists():
            excel_img_path = excel_dir / cand_img_name
            if excel_img_path.exists():
                try:
                    with PILImage.open(excel_img_path) as pimg:
                        orig_w, orig_h = pimg.size
                        max_w, max_h = 145, 95
                        scale = min(max_w / orig_w, max_h / orig_h, 1.0)
                        thumb_w_excel = int(orig_w * scale)
                        thumb_h_excel = int(orig_h * scale)

                    xl_img_excel = XLImage(str(excel_img_path))
                    xl_img_excel.width = thumb_w_excel
                    xl_img_excel.height = thumb_h_excel
                    ws.add_image(xl_img_excel, f"D{row_idx}")
                except Exception as img_err:
                    print(f"Warning embedding Excel image {cand_img_name}: {img_err}")
                    ws.cell(row=row_idx, column=4).value = cand_img_name

        max_thumb_h = max(thumb_h_pdf, thumb_h_excel)
        row_height = max(80, int(max_thumb_h * 0.75) + 15) if max_thumb_h > 0 else 40

        for c_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=c_idx)
            c.border = thin_border
            if c_idx in [1, 2, 3, 4, 6]:
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            if c_idx == 7:
                c.fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
        ws.row_dimensions[row_idx].height = row_height
        row_idx += 1

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 45
    ws.column_dimensions['F'].width = 32
    ws.column_dimensions['G'].width = 65


@app.get("/api/download-existing-alt-excel/{session_id}")
def download_existing_alt_excel(session_id: str, tab: Optional[str] = "pdf"):
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session expired or not found. Please upload or reload your PDF document.")

    actual_session_id = session_data.get("session_id", session_id)
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    orig_name = session_data.get("filename", "document.pdf")
    base_name = Path(orig_name).stem

    figures_dir = session_path / "figures"
    formulas_dir = session_path / "formulas"

    figures_list = session_data.get("figures", [])
    formulas_list = session_data.get("formulas", [])

    from openpyxl import Workbook

    wb = Workbook()

    # Sheet 1: PDF Figures
    ws_figures = wb.active
    ws_figures.title = "PDF Figures"
    populate_existing_alt_sheet(ws_figures, figures_list, figures_dir, label_prefix="PDF Figure")

    # Sheet 2: PDF Formulas
    ws_formulas = wb.create_sheet(title="PDF Formulas")
    populate_existing_alt_sheet(ws_formulas, formulas_list, formulas_dir, label_prefix="PDF Formula")

    # Set active sheet according to current view tab
    if tab == "formula":
        wb.active = ws_formulas
    else:
        wb.active = ws_figures

    excel_file_path = session_path / f"existing_alt_{base_name}.xlsx"
    save_openpyxl_workbook(wb, excel_file_path)

    return FileResponse(
        excel_file_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"existing_alt_{base_name}.xlsx"
    )

@app.get("/api/download-missing-alt-excel/{session_id}")
def download_missing_alt_excel(session_id: str, tab: Optional[str] = "pdf"):
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session expired or not found. Please upload or reload your PDF document.")

    actual_session_id = session_data.get("session_id", session_id)
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    orig_name = session_data.get("filename", "document.pdf")
    base_name = Path(orig_name).stem

    figures_dir = session_path / "figures"
    excel_dir = session_path / "excel_images"
    formulas_dir = session_path / "formulas"

    all_figures = session_data.get("figures", [])
    all_formulas = session_data.get("formulas", [])
    
    # Filter missing alt figures
    missing_figures = []
    for f in all_figures:
        has_effective_alt = bool(f.get("has_alt") or f.get("alt_text") or (f.get("excel_match") and f["excel_match"].get("alt_text")))
        if not has_effective_alt:
            missing_figures.append(f)

    # Filter missing alt formulas
    missing_formulas = []
    for f in all_formulas:
        has_effective_alt = bool(f.get("has_alt") or f.get("alt_text") or f.get("actual_text") or (f.get("excel_match") and f["excel_match"].get("alt_text")))
        if not has_effective_alt:
            missing_formulas.append(f)

    excel_records = session_data.get("excel_records", [])
    records_by_row = {r.get("row") or r.get("row_index"): r for r in excel_records if (r.get("row") or r.get("row_index"))}

    from openpyxl import Workbook

    wb = Workbook()

    # Sheet 1: Missing Alt Figures
    ws_figures = wb.active
    ws_figures.title = "Missing Alt Figures"
    populate_missing_alt_sheet(ws_figures, missing_figures, figures_dir, excel_dir, records_by_row, label_prefix="PDF Figure")

    # Sheet 2: Missing Alt Formulas
    ws_formulas = wb.create_sheet(title="Missing Alt Formulas")
    populate_missing_alt_sheet(ws_formulas, missing_formulas, formulas_dir, excel_dir, records_by_row, label_prefix="PDF Formula")

    if tab == "formula":
        wb.active = ws_formulas
    else:
        wb.active = ws_figures

    excel_file_path = session_path / f"missing_alt_{base_name}.xlsx"
    save_openpyxl_workbook(wb, excel_file_path)

    return FileResponse(
        excel_file_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"missing_alt_{base_name}.xlsx"
    )

@app.get("/api/download-missing-alt-zip/{session_id}")
def download_missing_alt_zip(session_id: str, tab: Optional[str] = "pdf"):
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session expired or not found. Please upload or reload your PDF document.")

    actual_session_id = session_data.get("session_id", session_id)
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    orig_name = session_data.get("filename", "document.pdf")
    base_name = Path(orig_name).stem

    figures_dir = session_path / "figures"
    excel_dir = session_path / "excel_images"
    all_figures = session_data.get("figures", [])
    
    # Filter missing alt figures (no PDF /Alt, no saved alt, no excel authoritative alt)
    missing_figures = []
    for f in all_figures:
        has_effective_alt = bool(f.get("has_alt") or f.get("alt_text") or (f.get("excel_match") and f["excel_match"].get("alt_text")))
        if not has_effective_alt:
            missing_figures.append(f)

    # Also check formulas
    formulas_dir = session_path / "formulas"
    all_formulas = session_data.get("formulas", [])
    missing_formulas = []
    for f in all_formulas:
        has_effective_alt = bool(f.get("has_alt") or f.get("alt_text") or f.get("actual_text") or (f.get("excel_match") and f["excel_match"].get("alt_text")))
        if not has_effective_alt:
            missing_formulas.append(f)

    is_formula_mode = (tab == "formula")
    target_items = missing_formulas if is_formula_mode else missing_figures
    target_dir = formulas_dir if is_formula_mode else figures_dir
    item_type_label = "Formulas" if is_formula_mode else "Figures"

    excel_records = session_data.get("excel_records", [])
    records_by_row = {r.get("row") or r.get("row_index"): r for r in excel_records if (r.get("row") or r.get("row_index"))}

    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zip_file:
        # CSV Manifest
        csv_buffer = io.StringIO()
        writer = csv.writer(csv_buffer)
        writer.writerow(["S.No", "Page No", "PDF Image File", "Excel Candidate Image", "Excel Alt Text", "Status", "Client Required Alt Text"])
        for seq_idx, item in enumerate(target_items, start=1):
            pdf_img_name = item.get("crop_filename", "")
            cand_match = item.get("candidate_match") or item.get("excel_match")
            best_row = None
            if not cand_match and item.get("match_debug") and item["match_debug"].get("best_single_row"):
                best_row = item["match_debug"]["best_single_row"]
                cand_match = records_by_row.get(best_row)
            elif cand_match:
                best_row = cand_match.get("row") or cand_match.get("row_index")

            cand_img_name = cand_match.get("image_filename") if cand_match else None
            cand_alt_val = (cand_match.get("alt_text") or cand_match.get("raw_alt_text") or "") if cand_match else ""
            conf_val = item.get("confidence") or (item.get("match_debug") and item["match_debug"].get("score")) or 0.0
            conf_pct = int(float(conf_val) * 100) if conf_val else 0

            if cand_match and best_row:
                status_val = f"Low Match - Verify Row {best_row} ({conf_pct}% match)"
            else:
                status_val = "No Match in Excel"

            if pdf_img_name and (target_dir / pdf_img_name).exists():
                folder_name = "missing_alt_formulas" if is_formula_mode else "missing_alt_images"
                zip_file.write(target_dir / pdf_img_name, arcname=f"{folder_name}/{pdf_img_name}")
            if cand_img_name and excel_dir.exists() and (excel_dir / cand_img_name).exists():
                zip_file.write(excel_dir / cand_img_name, arcname=f"excel_candidates/{cand_img_name}")

            writer.writerow([
                seq_idx,
                item.get("page_number", ""),
                pdf_img_name,
                cand_img_name or "None",
                cand_alt_val,
                status_val,
                ""
            ])
        zip_file.writestr(f"missing_{item_type_label.lower()}_manifest.csv", csv_buffer.getvalue())

        # Formatted Excel Workbook template for client with embedded visual images (Multi-sheet: Figures & Formulas)
        try:
            from openpyxl import Workbook

            wb = Workbook()

            # Sheet 1: Missing Alt Figures
            ws_figures = wb.active
            ws_figures.title = "Missing Alt Figures"
            populate_missing_alt_sheet(ws_figures, missing_figures, figures_dir, excel_dir, records_by_row, label_prefix="PDF Figure")

            # Sheet 2: Missing Alt Formulas
            ws_formulas = wb.create_sheet(title="Missing Alt Formulas")
            populate_missing_alt_sheet(ws_formulas, missing_formulas, formulas_dir, excel_dir, records_by_row, label_prefix="PDF Formula")

            if tab == "formula":
                wb.active = ws_formulas
            else:
                wb.active = ws_figures

            excel_buf = io.BytesIO()
            wb.save(excel_buf)
            zip_file.writestr(f"missing_alt_template_{base_name}.xlsx", excel_buf.getvalue())
        except Exception as err:
            print(f"Excel template generation warning: {err}")

    download_filename = f"missing_alt_{item_type_label.lower()}_{base_name}.zip"
    return FileResponse(
        zip_path,
        media_type="application/zip",
        filename=download_filename
    )

@app.get("/api/download-figures-zip/{session_id}")
def download_figures_zip(session_id: str):
    return download_zip(session_id)

@app.get("/api/download-zip/{session_id}")
def download_zip(session_id: str):
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session figures not found")

    actual_session_id = session_data.get("session_id", session_id)
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    figures_dir = session_path / "figures"
    if not figures_dir.exists():
        raise HTTPException(status_code=404, detail="Session figures not found")

    zip_path = session_path / f"pdf_figures_{actual_session_id[:8]}.zip"
    
    # Generate CSV manifest inside ZIP
    csv_buffer = io.StringIO()
    writer = csv.writer(csv_buffer)
    writer.writerow(["Figure ID", "Page", "MCIDs", "Status", "Alt Text", "BBox [x0, y0, x1, y1]", "Width (pt)", "Height (pt)", "Image File"])

    figures_meta = session_data.get("figures", [])
    meta_by_file = {f.get("crop_filename"): f for f in figures_meta}

    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for img_file in figures_dir.glob("*.png"):
            zip_file.write(img_file, arcname=f"figures/{img_file.name}")
            meta = meta_by_file.get(img_file.name, {})
            writer.writerow([
                meta.get("figure_id", ""),
                meta.get("page_number", ""),
                str(meta.get("mcids", "")),
                meta.get("status", ""),
                meta.get("alt_text", "") or "",
                str(meta.get("bbox", "")),
                meta.get("bbox_width", ""),
                meta.get("bbox_height", ""),
                img_file.name
            ])

        zip_file.writestr("figures_manifest.csv", csv_buffer.getvalue())

    return FileResponse(
        zip_path,
        media_type="application/zip",
        filename=f"extracted_figures_{actual_session_id[:8]}.zip"
    )

@app.get("/api/download-json/{session_id}")
def download_json(session_id: str):
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session not found")
    return JSONResponse(content=session_data)

class SingleAltInjection(BaseModel):
    figure_id: int
    alt_text: str

class BatchAltInjection(BaseModel):
    injections: Optional[Dict[int, str]] = None

class BatchAltRemoval(BaseModel):
    figure_ids: Optional[List[int]] = None

class SingleFormulaAltInjection(BaseModel):
    formula_id: int
    alt_text: str

class BatchFormulaAltInjection(BaseModel):
    injections: Optional[Dict[int, str]] = None

class BatchFormulaAltRemoval(BaseModel):
    formula_ids: Optional[List[int]] = None

class UnselectedFiguresDownload(BaseModel):
    figure_ids: List[int]

class UpdateItemAltRequest(BaseModel):
    item_type: str  # 'formula' or 'figure'
    item_id: int
    alt_text: str

@app.post("/api/update-item-alt/{session_id}")
async def update_item_alt(session_id: str, payload: UpdateItemAltRequest):
    """
    Updates and saves custom/edited ALT text for a specific figure or formula in the session cache.
    Persists the edited Alt text across views and exports.
    """
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session not found")

    actual_session_id = session_data.get("session_id", session_id)
    alt_clean = payload.alt_text.strip()
    has_alt = bool(alt_clean)

    if payload.item_type == "formula":
        formulas = session_data.get("formulas", [])
        for form in formulas:
            if form.get("formula_id") == payload.item_id:
                form["alt_text"] = alt_clean
                form["has_alt"] = has_alt
                if not form.get("status_label") or form["status_label"] in ["No Match", "Missing /Alt", "Missing Alt"]:
                    form["status_label"] = "Has Alt" if has_alt else "Missing Alt"
                break
        session_data["formulas"] = formulas
        session_data["has_formula_alt_count"] = sum(1 for f in formulas if f.get("has_alt") or f.get("excel_match"))
        session_data["missing_formula_alt_count"] = len(formulas) - session_data["has_formula_alt_count"]
    else:
        figures = session_data.get("figures", [])
        for fig in figures:
            if fig.get("figure_id") == payload.item_id:
                fig["alt_text"] = alt_clean
                fig["has_alt"] = has_alt
                if not fig.get("status_label") or fig["status_label"] in ["No Match", "Missing /Alt", "Missing Alt"]:
                    fig["status_label"] = "Has Alt" if has_alt else "Missing Alt"
                break
        session_data["figures"] = figures
        session_data["has_alt_count"] = sum(1 for f in figures if f.get("has_alt") or f.get("excel_match"))
        session_data["missing_alt_count"] = len(figures) - session_data["has_alt_count"]

    session_cache[actual_session_id] = session_data
    session_cache[session_id] = session_data
    save_session_to_disk(actual_session_id, session_data)

    return JSONResponse(content={
        "status": "success",
        "session_id": actual_session_id,
        "item_type": payload.item_type,
        "item_id": payload.item_id,
        "alt_text": alt_clean,
        "has_alt": has_alt
    })

@app.post("/api/inject-alt/{session_id}")
async def inject_all_alt(session_id: str, payload: Optional[BatchAltInjection] = None):
    """
    Injects alt texts into the PDF's StructTreeRoot elements where /S == '/Figure'.
    Uses either the provided injections map or defaults to all matched Excel alt texts.
    Saves the new accessible PDF to injected_accessible.pdf and returns updated metrics.
    """
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session not found")

    actual_session_id = session_data.get("session_id", session_id)
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    base_pdf = session_path / "injected_accessible.pdf"
    if not base_pdf.exists():
        base_pdf = session_path / "input.pdf"

    if not base_pdf.exists():
        raise HTTPException(status_code=404, detail="Input PDF not found for session")

    temp_output = session_path / "injected_temp.pdf"
    target_output = session_path / "injected_accessible.pdf"
    figures = session_data.get("figures", [])

    # Build injection dictionary (1-indexed figure_id -> alt_text)
    injections_map = {}
    if payload and payload.injections:
        injections_map = {int(k): str(v).strip() for k, v in payload.injections.items() if not is_duplicate_marker(str(v))}
    else:
        for f in figures:
            fig_id = f.get("figure_id")
            # Use matched Excel alt text or existing alt text
            alt = (f.get("excel_match") and f["excel_match"].get("alt_text")) or f.get("alt_text")
            if alt and fig_id and not is_duplicate_marker(alt):
                injections_map[fig_id] = alt.strip()

    if not injections_map:
        raise HTTPException(status_code=400, detail="No alt texts found to inject.")

    # Run injection via FigureExtractor
    extractor = None
    try:
        extractor = FigureExtractor(str(base_pdf))
        count = extractor.inject_alt_texts(injections_map, str(temp_output))
    finally:
        if extractor:
            extractor.close()

    if temp_output.exists():
        shutil.move(str(temp_output), str(target_output))

    # Update session cache
    for f in figures:
        fid = f.get("figure_id")
        if fid in injections_map:
            f["alt_text"] = injections_map[fid]
            f["has_alt"] = True
            f["status_label"] = "Injected"

    has_alt_count = sum(1 for f in figures if f["has_alt"])
    missing_alt_count = len(figures) - has_alt_count

    session_data["figures"] = figures
    session_data["has_alt_count"] = has_alt_count
    session_data["missing_alt_count"] = missing_alt_count
    session_data["has_injected_pdf"] = True

    session_cache[actual_session_id] = session_data
    session_cache[session_id] = session_data
    save_session_to_disk(actual_session_id, session_data)

    return {
        "session_id": actual_session_id,
        "status": "success",
        "injected_count": count,
        "total_figures": len(figures),
        "has_alt_count": has_alt_count,
        "missing_alt_count": missing_alt_count,
        "download_url": f"/api/download-injected-pdf/{actual_session_id}",
        "figures": figures
    }

@app.post("/api/inject-single-alt/{session_id}")
async def inject_single_alt(session_id: str, payload: SingleAltInjection):
    """
    Injects alt text for a single figure into the PDF.
    """
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session not found")

    actual_session_id = session_data.get("session_id", session_id)
    alt_clean = payload.alt_text.strip()
    if is_duplicate_marker(alt_clean):
        raise HTTPException(status_code=400, detail="Cannot inject 'Duplicate' placeholder as ALT text.")

    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    base_pdf = session_path / "injected_accessible.pdf"
    if not base_pdf.exists():
        base_pdf = session_path / "input.pdf"

    if not base_pdf.exists():
        raise HTTPException(status_code=404, detail="Base PDF not found")

    temp_output = session_path / "injected_temp.pdf"
    target_output = session_path / "injected_accessible.pdf"

    extractor = None
    try:
        extractor = FigureExtractor(str(base_pdf))
        count = extractor.inject_alt_texts({payload.figure_id: payload.alt_text.strip()}, str(temp_output))
    finally:
        if extractor:
            extractor.close()

    if temp_output.exists():
        shutil.move(str(temp_output), str(target_output))

    # Update session cache
    figures = session_data.get("figures", [])
    for f in figures:
        if f.get("figure_id") == payload.figure_id:
            f["alt_text"] = payload.alt_text.strip()
            f["has_alt"] = True
            f["status_label"] = "Injected"

    has_alt_count = sum(1 for f in figures if f["has_alt"])
    missing_alt_count = len(figures) - has_alt_count
    session_data["figures"] = figures
    session_data["has_alt_count"] = has_alt_count
    session_data["missing_alt_count"] = missing_alt_count
    session_data["has_injected_pdf"] = True

    session_cache[actual_session_id] = session_data
    session_cache[session_id] = session_data
    save_session_to_disk(actual_session_id, session_data)

    return {
        "session_id": actual_session_id,
        "status": "success",
        "figure_id": payload.figure_id,
        "injected_alt": payload.alt_text.strip(),
        "has_alt_count": has_alt_count,
        "missing_alt_count": missing_alt_count,
        "download_url": f"/api/download-injected-pdf/{actual_session_id}"
    }

@app.post("/api/remove-alt/{session_id}")
async def remove_alt(session_id: str, payload: Optional[BatchAltRemoval] = None):
    """
    Removes /Alt accessibility texts from the PDF's StructTreeRoot elements where /S == '/Figure'.
    Can remove for selected figure IDs or for all figures if none specified.
    Saves the updated accessible PDF and returns updated metrics.
    """
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session not found")

    actual_session_id = session_data.get("session_id", session_id)
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    base_pdf = session_path / "injected_accessible.pdf"
    if not base_pdf.exists():
        base_pdf = session_path / "input.pdf"

    if not base_pdf.exists():
        raise HTTPException(status_code=404, detail="Base PDF not found for session")

    temp_output = session_path / "injected_temp.pdf"
    target_output = session_path / "injected_accessible.pdf"
    figures = session_data.get("figures", [])

    target_ids = payload.figure_ids if (payload and payload.figure_ids) else None

    extractor = None
    try:
        extractor = FigureExtractor(str(base_pdf))
        count = extractor.remove_alt_texts(target_ids, str(temp_output))
    finally:
        if extractor:
            extractor.close()

    if temp_output.exists():
        shutil.move(str(temp_output), str(target_output))

    target_set = set(target_ids) if target_ids else None
    for f in figures:
        fid = f.get("figure_id")
        if target_set is None or fid in target_set:
            f["alt_text"] = ""
            f["has_alt"] = False
            f["status_label"] = "Missing /Alt"

    has_alt_count = sum(1 for f in figures if f["has_alt"])
    missing_alt_count = len(figures) - has_alt_count

    session_data["figures"] = figures
    session_data["has_alt_count"] = has_alt_count
    session_data["missing_alt_count"] = missing_alt_count
    session_data["has_injected_pdf"] = True

    session_cache[actual_session_id] = session_data
    session_cache[session_id] = session_data
    save_session_to_disk(actual_session_id, session_data)

    return {
        "session_id": actual_session_id,
        "status": "success",
        "removed_count": count,
        "total_figures": len(figures),
        "has_alt_count": has_alt_count,
        "missing_alt_count": missing_alt_count,
        "download_url": f"/api/download-injected-pdf/{actual_session_id}",
        "figures": figures
    }

@app.post("/api/inject-formula-alt/{session_id}")
async def inject_formula_alt(session_id: str, payload: Optional[BatchFormulaAltInjection] = None):
    """
    Injects alt texts into the PDF's StructTreeRoot elements where /S is a formula tag.
    Uses either provided injections map or defaults to all matched Excel formula alt texts.
    """
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session not found")

    actual_session_id = session_data.get("session_id", session_id)
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    base_pdf = session_path / "injected_accessible.pdf"
    if not base_pdf.exists():
        base_pdf = session_path / "input.pdf"

    if not base_pdf.exists():
        raise HTTPException(status_code=404, detail="Input PDF not found for session")

    temp_output = session_path / "injected_temp.pdf"
    target_output = session_path / "injected_accessible.pdf"
    formulas = session_data.get("formulas", [])

    injections_map = {}
    if payload and payload.injections:
        injections_map = {int(k): str(v).strip() for k, v in payload.injections.items() if not is_duplicate_marker(str(v))}
    else:
        for form in formulas:
            fid = form.get("formula_id")
            alt = (form.get("excel_match") and form["excel_match"].get("alt_text")) or form.get("alt_text") or form.get("actual_text")
            if alt and fid and not is_duplicate_marker(alt):
                injections_map[fid] = alt.strip()

    if not injections_map:
        raise HTTPException(status_code=400, detail="No formula alt texts provided or available to inject.")

    extractor = None
    try:
        extractor = FigureExtractor(str(base_pdf))
        count = extractor.inject_formula_alt_texts(injections_map, str(temp_output))
    finally:
        if extractor:
            extractor.close()

    if temp_output.exists():
        shutil.move(str(temp_output), str(target_output))

    for form in formulas:
        fid = form.get("formula_id")
        if fid in injections_map:
            form["alt_text"] = injections_map[fid]
            form["has_alt"] = True
            form["status_label"] = "Injected"

    has_formula_alt_count = sum(1 for f in formulas if f["has_alt"])
    missing_formula_alt_count = len(formulas) - has_formula_alt_count

    session_data["formulas"] = formulas
    session_data["has_formula_alt_count"] = has_formula_alt_count
    session_data["missing_formula_alt_count"] = missing_formula_alt_count
    session_data["has_injected_pdf"] = True

    session_cache[actual_session_id] = session_data
    session_cache[session_id] = session_data
    save_session_to_disk(actual_session_id, session_data)

    return {
        "session_id": actual_session_id,
        "status": "success",
        "injected_count": count,
        "total_formulas": len(formulas),
        "has_formula_alt_count": has_formula_alt_count,
        "missing_formula_alt_count": missing_formula_alt_count,
        "download_url": f"/api/download-injected-pdf/{actual_session_id}",
        "formulas": formulas
    }

@app.post("/api/inject-single-formula-alt/{session_id}")
async def inject_single_formula_alt(session_id: str, payload: SingleFormulaAltInjection):
    """
    Injects alt text for a single formula into the PDF.
    """
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session not found")

    actual_session_id = session_data.get("session_id", session_id)
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    base_pdf = session_path / "injected_accessible.pdf"
    if not base_pdf.exists():
        base_pdf = session_path / "input.pdf"

    if not base_pdf.exists():
        raise HTTPException(status_code=404, detail="Base PDF not found")

    temp_output = session_path / "injected_temp.pdf"
    target_output = session_path / "injected_accessible.pdf"

    formulas = session_data.get("formulas", [])
    alt_to_inject = payload.alt_text.strip()
    if is_duplicate_marker(alt_to_inject):
        raise HTTPException(status_code=400, detail="Cannot inject 'Duplicate' placeholder as formula ALT text.")

    if not alt_to_inject:
        for f in formulas:
            if f.get("formula_id") == payload.formula_id:
                cand = ((f.get("excel_match") and f["excel_match"].get("alt_text")) or f.get("alt_text") or f.get("actual_text") or "").strip()
                if cand and not is_duplicate_marker(cand):
                    alt_to_inject = cand
                break

    if not alt_to_inject:
        raise HTTPException(status_code=400, detail="No alt text available to inject for formula.")

    extractor = None
    try:
        extractor = FigureExtractor(str(base_pdf))
        count = extractor.inject_formula_alt_texts({payload.formula_id: alt_to_inject}, str(temp_output))
    finally:
        if extractor:
            extractor.close()

    if temp_output.exists():
        shutil.move(str(temp_output), str(target_output))

    for f in formulas:
        if f.get("formula_id") == payload.formula_id:
            f["alt_text"] = alt_to_inject
            f["has_alt"] = True
            f["status_label"] = "Injected"

    has_formula_alt_count = sum(1 for f in formulas if f["has_alt"])
    missing_formula_alt_count = len(formulas) - has_formula_alt_count

    session_data["formulas"] = formulas
    session_data["has_formula_alt_count"] = has_formula_alt_count
    session_data["missing_formula_alt_count"] = missing_formula_alt_count
    session_data["has_injected_pdf"] = True

    session_cache[actual_session_id] = session_data
    session_cache[session_id] = session_data
    save_session_to_disk(actual_session_id, session_data)

    return {
        "session_id": actual_session_id,
        "status": "success",
        "formula_id": payload.formula_id,
        "injected_alt": alt_to_inject,
        "has_formula_alt_count": has_formula_alt_count,
        "missing_formula_alt_count": missing_formula_alt_count,
        "download_url": f"/api/download-injected-pdf/{actual_session_id}"
    }

@app.post("/api/remove-formula-alt/{session_id}")
async def remove_formula_alt(session_id: str, payload: Optional[BatchFormulaAltRemoval] = None):
    """
    Removes /Alt and /ActualText from the PDF's formula tags.
    """
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session not found")

    actual_session_id = session_data.get("session_id", session_id)
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    base_pdf = session_path / "injected_accessible.pdf"
    if not base_pdf.exists():
        base_pdf = session_path / "input.pdf"

    if not base_pdf.exists():
        raise HTTPException(status_code=404, detail="Base PDF not found for session")

    temp_output = session_path / "injected_temp.pdf"
    target_output = session_path / "injected_accessible.pdf"
    formulas = session_data.get("formulas", [])

    target_ids = payload.formula_ids if (payload and payload.formula_ids) else None

    extractor = None
    try:
        extractor = FigureExtractor(str(base_pdf))
        count = extractor.remove_formula_alt_texts(target_ids, str(temp_output))
    finally:
        if extractor:
            extractor.close()

    if temp_output.exists():
        shutil.move(str(temp_output), str(target_output))

    target_set = set(target_ids) if target_ids else None
    for f in formulas:
        fid = f.get("formula_id")
        if target_set is None or fid in target_set:
            f["alt_text"] = ""
            f["actual_text"] = ""
            f["has_alt"] = False
            f["status"] = "Missing Alt"

    has_formula_alt_count = sum(1 for f in formulas if f["has_alt"])
    missing_formula_alt_count = len(formulas) - has_formula_alt_count

    session_data["formulas"] = formulas
    session_data["has_formula_alt_count"] = has_formula_alt_count
    session_data["missing_formula_alt_count"] = missing_formula_alt_count
    session_data["has_injected_pdf"] = True

    session_cache[actual_session_id] = session_data
    session_cache[session_id] = session_data
    save_session_to_disk(actual_session_id, session_data)

    return {
        "session_id": actual_session_id,
        "status": "success",
        "removed_count": count,
        "total_formulas": len(formulas),
        "has_formula_alt_count": has_formula_alt_count,
        "missing_formula_alt_count": missing_formula_alt_count,
        "download_url": f"/api/download-injected-pdf/{actual_session_id}",
        "formulas": formulas
    }

@app.get("/api/download-injected-pdf/{session_id}")
def download_injected_pdf(session_id: str):
    session_data = get_session_data(session_id)
    actual_session_id = session_data.get("session_id", session_id) if session_data else session_id
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    injected_pdf = session_path / "injected_accessible.pdf"
    if not injected_pdf.exists():
        raise HTTPException(status_code=404, detail="Injected accessible PDF not found for this session")

    orig_name = (session_data or {}).get("filename", "document.pdf")
    base_name = Path(orig_name).stem
    download_filename = f"{base_name}_with_injected_alt.pdf"

    return FileResponse(
        injected_pdf,
        media_type="application/pdf",
        filename=download_filename
    )

@app.post("/api/download-unselected-excel/{session_id}")
async def download_unselected_excel(session_id: str, payload: UnselectedFiguresDownload):
    session_data = get_session_data(session_id)
    if not session_data:
        raise HTTPException(status_code=404, detail="Session not found")

    actual_session_id = session_data.get("session_id", session_id)
    session_path = resolve_session_path(session_id) or (SESSIONS_DIR / actual_session_id)
    figures_dir = session_path / "figures"
    if not figures_dir.exists():
        raise HTTPException(status_code=404, detail="Session figures not found")

    unselected_ids = set(payload.figure_ids)
    all_figures = session_data.get("figures", [])

    unselected_figures = [f for f in all_figures if f.get("figure_id") in unselected_ids]
    if not unselected_figures:
        raise HTTPException(status_code=400, detail="No unselected figures found")

    excel_file_path = session_path / f"unselected_figures_{actual_session_id[:8]}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Unselected Figures"

    ws.append(["Figure ID", "Page", "MCIDs", "Status", "Alt Text", "BBox [x0, y0, x1, y1]", "Width (pt)", "Height (pt)", "Image File"])

    row_idx = 2
    for fig in unselected_figures:
        ws.append([
            sanitize_excel_val(fig.get("figure_id", "")),
            sanitize_excel_val(fig.get("page_number", "")),
            sanitize_excel_val(str(fig.get("mcids", ""))),
            sanitize_excel_val(fig.get("status", "")),
            sanitize_excel_val(fig.get("alt_text", "") or ""),
            sanitize_excel_val(str(fig.get("bbox", ""))),
            sanitize_excel_val(fig.get("bbox_width", "")),
            sanitize_excel_val(fig.get("bbox_height", "")),
            sanitize_excel_val(fig.get("crop_filename", ""))
        ])

        img_filename = fig.get("crop_filename")
        if img_filename:
            img_path = figures_dir / img_filename
            if img_path.exists():
                try:
                    excel_img = XLImage(str(img_path))
                    excel_img.width = 150
                    excel_img.height = 150
                    ws.add_image(excel_img, f"I{row_idx}")
                except Exception as e:
                    print(f"Warning: Could not embed image {img_filename}: {e}")

        row_idx += 1

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 8
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 25
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 12
    ws.column_dimensions['I'].width = 18

    save_openpyxl_workbook(wb, excel_file_path)

    return FileResponse(
        excel_file_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"unselected_figures_{session_id[:8]}.xlsx"
    )

# Mount static frontend
app.mount("/", StaticFiles(directory=str(STATIC_DIR), html=True), name="static")

import socket

def find_available_port(start_port=8000, max_attempts=20):
    for p in range(start_port, start_port + max_attempts):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            try:
                s.bind(('127.0.0.1', p))
                return p
            except OSError:
                continue
    return start_port

if __name__ == "__main__":
    import uvicorn
    port = find_available_port(8000)
    url = f"http://127.0.0.1:{port}"

    print("\n" + "=" * 60)
    print("  PDF ALT-TEXT & FIGURE EXTRACTION TOOL")
    print(f"  Server running at: {url}")
    print("  Keep this window open while using the application.")
    print("  Close this window when you are done to exit.")
    print("=" * 60 + "\n")
    
    def open_browser():
        try:
            webbrowser.open(url)
        except Exception as e:
            print(f"Could not automatically open browser: {e}")

    # Automatically open browser after 1.5 seconds once server starts
    threading.Timer(1.5, open_browser).start()
    uvicorn.run(app, host="127.0.0.1", port=port, log_config=None)


